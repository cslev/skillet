#!/usr/bin/env python3
"""Helpers for building and editing decks with python-pptx (pitch-deck skill).

Import from a build script:

    import sys; sys.path.insert(0, "<skill-dir>")
    from deck_lib import *

Everything here is brand-neutral: pass colours and fonts in (from the template
analysis or the user's answers). Nothing is generated from an outline; the build
script written for each run decides content and layout.

Covers the traps met in practice: template demo slides to drop, placeholders by
`idx`, run formatting lost on `text =`, unfilled "Click to add text" placeholders,
native charts with a manual plot layout (so annotations can be positioned exactly),
the Thank-you slide, hidden slides, and a before/after guard for single-slide edits.
"""

import copy
import hashlib
import re

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

EMU = 914400


def rgb(hexstr):
    h = hexstr.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


# ---------------------------------------------------------------- opening

def open_template(path, keep_slides=False):
    """Open a template as the base deck. Drops its demo slides (masters and
    layouts stay) unless keep_slides. Returns the Presentation."""
    prs = Presentation(path)
    if not keep_slides:
        ids = prs.slides._sldIdLst
        for s in list(ids):
            prs.part.drop_rel(s.rId)
            ids.remove(s)
    return prs


def blank_deck(width_in=13.333, height_in=7.5):
    """A 16:9 deck with python-pptx's default master, for the no-template case."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(width_in), Inches(height_in)
    return prs


def layouts(prs):
    return {l.name: l for l in prs.slide_layouts}


def add_slide(prs, layout_name):
    return prs.slides.add_slide(layouts(prs)[layout_name])


def find_closing_layout(prs):
    """The template's own thank-you / closing layout, or None."""
    for l in prs.slide_layouts:
        if re.search(r"\bclosing\b|thank\s*you|\bend\s*slide\b", l.name, re.I):
            return l
    return None


def theme_info(prs):
    """Theme fonts and colours: {'major', 'minor', 'colors': {name: '#RRGGBB'}}."""
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    theme = prs.slide_masters[0].part.part_related_by(RT.THEME)
    x = etree.fromstring(theme.blob)
    ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    out = {"colors": {}}
    for tag in ("major", "minor"):
        e = x.find(f".//a:{tag}Font/a:latin", ns)
        out[tag] = e.get("typeface") if e is not None else None
    cs = x.find(".//a:clrScheme", ns)
    for c in cs:
        name = etree.QName(c).localname
        v = c[0].get("val") or c[0].get("lastClr")
        out["colors"][name] = "#" + v
    return out


# ------------------------------------------------------------ placeholders

def ph(slide, idx):
    for p in slide.placeholders:
        if p.placeholder_format.idx == idx:
            return p
    return None


def drop_ph(slide, idx):
    p = ph(slide, idx)
    if p is not None:
        p._element.getparent().remove(p._element)


def has_own_fill(shape):
    """True for a placeholder that carries its own solid fill (a decorative panel)."""
    return shape._element.find(".//" + qn("p:spPr") + "/" + qn("a:solidFill")) is not None


def scrub_placeholders(slide, keep=()):
    """Delete placeholders with no text (they render 'Click to add text'),
    except idx in `keep` and filled decorative panels."""
    for p in list(slide.placeholders):
        if p.placeholder_format.idx in keep or has_own_fill(p):
            continue
        if p.has_text_frame and not p.text_frame.text.strip():
            p._element.getparent().remove(p._element)


def set_text(shape, text, size=None):
    """Replace a shape's text, keeping the first run's formatting (setting
    `text_frame.text` on a placeholder with runs drops it). `\\n` makes paragraphs."""
    tf = shape.text_frame
    first = tf.paragraphs[0]
    rpr = None
    if first.runs:
        r = first.runs[0]._r.find(qn("a:rPr"))
        rpr = copy.deepcopy(r) if r is not None else None
    ppr = first._p.find(qn("a:pPr"))
    ppr = copy.deepcopy(ppr) if ppr is not None else None
    txbody = tf._txBody
    for p in txbody.findall(qn("a:p")):
        txbody.remove(p)
    for line in text.split("\n"):
        p = etree.SubElement(txbody, qn("a:p"))
        if ppr is not None:
            p.append(copy.deepcopy(ppr))
        r = etree.SubElement(p, qn("a:r"))
        if rpr is not None:
            r.append(copy.deepcopy(rpr))
        t = etree.SubElement(r, qn("a:t"))
        t.text = line
    if size:
        for par in tf.paragraphs:
            for run in par.runs:
                run.font.size = Pt(size)


def set_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def finish(slide, notes, source=None, source_idx=None, drop_idx=()):
    """Notes, the source line (placeholder `source_idx` if the layout has one),
    and removal of footers listed in drop_idx."""
    for i in drop_idx:
        drop_ph(slide, i)
    if source_idx is not None:
        if source:
            set_text(ph(slide, source_idx), source)
        else:
            drop_ph(slide, source_idx)
    set_notes(slide, notes)


def clear_slide(slide, keep_placeholders=True):
    """Remove everything the previous build added, keeping placeholders
    (title, footer, source) so the layout stays intact. For redoing one slide."""
    for sh in list(slide.shapes):
        if keep_placeholders and sh.is_placeholder:
            continue
        sh._element.getparent().remove(sh._element)


# ------------------------------------------------------------------ shapes

def _bullet(par, char="•", indent=0.22):
    pPr = par._p.get_or_add_pPr()
    pPr.set("marL", str(int(Inches(indent))))
    pPr.set("indent", str(-int(Inches(indent))))
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum"):
        for e in pPr.findall(qn(tag)):
            pPr.remove(e)
    etree.SubElement(pPr, qn("a:buChar")).set("char", char)


def textbox(slide, x, y, w, h, paras, font, size=18, color="#1F2328", bullets=False,
            align=PP_ALIGN.LEFT, space=6, anchor=MSO_ANCHOR.TOP):
    """Inches. paras: list of str or (str, {size,color,bold,bullet,space,align,font})."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    for i, item in enumerate(paras):
        text, o = (item, {}) if isinstance(item, str) else item
        par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        par.alignment = o.get("align", align)
        par.space_after = Pt(o.get("space", space))
        par.line_spacing = 1.1
        r = par.add_run()
        r.text = text
        r.font.name = o.get("font", font)
        r.font.size = Pt(o.get("size", size))
        r.font.color.rgb = rgb(o.get("color", color))
        r.font.bold = o.get("bold", False)
        if o.get("bullet", bullets):
            _bullet(par)
    return tb


def outline_box(slide, x, y, w, h, color="#57606A", rounded=False, weight=1):
    """Outline-only rectangle: no fill, no theme style (which would add shadow,
    blue outline, white text)."""
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
                                 Inches(x), Inches(y), Inches(w), Inches(h))
    st = shp._element.find(qn("p:style"))
    if st is not None:
        shp._element.remove(st)
    if rounded:
        shp.adjustments[0] = 0.08
    shp.fill.background()
    shp.line.color.rgb = rgb(color)
    shp.line.width = Pt(weight)
    shp.shadow.inherit = False
    return shp


def placeholder_box(slide, x, y, w, h, font, label="[ IMAGE PLACEHOLDER ]", desc="", color="#57606A"):
    """The standard image / diagram placeholder: outline, centred label, one-line brief."""
    shp = outline_box(slide, x, y, w, h, color)
    tf = shp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0.15)
    for i, (txt, sz) in enumerate([(label, 16), (desc, 14)]):
        if i and not txt:
            break
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = txt
        r.font.name = font
        r.font.size = Pt(sz)
        r.font.color.rgb = rgb(color)
    return shp


def add_table(slide, rows, x, y, w, col_widths, font, size=14, row_h=0.45):
    """Plain native table (first row is the header). Column widths in inches; they
    are set per column because resizing the frame does not rescale them."""
    gf = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(row_h * len(rows)))
    t = gf.table
    for i, cw in enumerate(col_widths):
        t.columns[i].width = Inches(cw)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            c = t.cell(ri, ci)
            c.text = str(val)
            for r in c.text_frame.paragraphs[0].runs:
                r.font.size = Pt(size)
                r.font.name = font
    return gf


def add_picture_fit(slide, path, x, y, w, h):
    """Place an image inside the box, keeping its aspect ratio, centred."""
    from PIL import Image
    iw, ih = Image.open(path).size
    scale = min(w / iw, h / ih)
    pw, phh = iw * scale, ih * scale
    return slide.shapes.add_picture(path, Inches(x + (w - pw) / 2), Inches(y + (h - phh) / 2), Inches(pw), Inches(phh))


# ------------------------------------------------------------------ charts

class ChartGeom:
    """Where things are on a chart with a manual inner plot layout (inches)."""

    def __init__(self, x, y, w, h, px, py, pw, ph, ncat, vmax):
        self.x, self.y, self.w, self.h = x, y, w, h
        self.px, self.py, self.pw, self.ph, self.ncat, self.vmax = px, py, pw, ph, ncat, vmax

    def slot_w(self):
        return self.pw * self.w / self.ncat

    def cat_center_x(self, i):
        return self.x + self.px * self.w + self.slot_w() * (i + 0.5)

    def cat_left_x(self, i):
        return self.x + self.px * self.w + self.slot_w() * i

    def value_y(self, v):
        return self.y + self.py * self.h + self.ph * self.h * (1 - v / self.vmax)

    def baseline_y(self):
        return self.value_y(0)


def _plot_layout(chart, px, py, pw, ph):
    pa = chart._chartSpace.find(qn("c:chart")).find(qn("c:plotArea"))
    for e in pa.findall(qn("c:layout")):
        pa.remove(e)
    pa.insert(0, etree.fromstring(
        '<c:layout xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart"><c:manualLayout>'
        '<c:layoutTarget val="inner"/><c:xMode val="edge"/><c:yMode val="edge"/>'
        f'<c:x val="{px}"/><c:y val="{py}"/><c:w val="{pw}"/><c:h val="{ph}"/></c:manualLayout></c:layout>'))


def _legend_layout(chart, x=0.12, y=0.0, w=0.76, h=0.07):
    """Pin the legend to one centred row. Wide w keeps it from wrapping; reduce
    entry spacing (shorter names), never the font size."""
    lg = chart._chartSpace.find(qn("c:chart")).find(qn("c:legend"))
    for e in lg.findall(qn("c:layout")):
        lg.remove(e)
    lg.find(qn("c:legendPos")).addnext(etree.fromstring(
        '<c:layout xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart"><c:manualLayout>'
        '<c:xMode val="edge"/><c:yMode val="edge"/>'
        f'<c:x val="{x}"/><c:y val="{y}"/><c:w val="{w}"/><c:h val="{h}"/></c:manualLayout></c:layout>'))


def bar_chart(slide, x, y, w, h, cats, series, colors, font, *, text="#1F2328", axis="#9AA4B0",
              number_format='0.0"%"', legend=False, vmax=100, gap=60, overlap=None,
              label_size=11, plot=(0.03, 0.14, 0.94, 0.70), point_colors=None):
    """Native clustered column chart, values labelled, value axis hidden.
    series: [(name, [values])]; a None value leaves a gap. plot=(x, y, w, h) as
    fractions of the frame (manual inner layout). Returns (chart, ChartGeom)."""
    cd = CategoryChartData()
    cd.categories = cats
    for name, vals in series:
        cd.add_series(name, vals)
    gf = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(x), Inches(y), Inches(w), Inches(h), cd)
    ch = gf.chart
    ch.font.name, ch.font.size, ch.font.color.rgb = font, Pt(14), rgb(text)
    ch.has_title = False
    ch.has_legend = legend
    if legend:
        ch.legend.position = XL_LEGEND_POSITION.TOP
        ch.legend.include_in_layout = False
        ch.legend.font.size, ch.legend.font.name = Pt(12), font
    va = ch.value_axis
    va.maximum_scale, va.minimum_scale = vmax, 0
    va.visible = False
    va.has_major_gridlines = False
    ca = ch.category_axis
    ca.tick_labels.font.size, ca.tick_labels.font.name = Pt(12), font
    ca.format.line.color.rgb = rgb(axis)
    ca.has_major_gridlines = False
    pl = ch.plots[0]
    pl.gap_width = gap
    pl.overlap = overlap if overlap is not None else (-5 if len(series) > 1 else 0)
    pl.has_data_labels = True
    d = pl.data_labels
    d.number_format, d.number_format_is_linked = number_format, False
    d.position = XL_LABEL_POSITION.OUTSIDE_END
    d.font.size, d.font.name, d.font.color.rgb = Pt(label_size), font, rgb(text)
    for s, c in zip(pl.series, colors):
        s.format.fill.solid()
        s.format.fill.fore_color.rgb = rgb(c)
        s.invert_if_negative = False
    for i, c in (point_colors or {}).items():
        pt = pl.series[0].points[i]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = rgb(c)
    _plot_layout(ch, *plot)
    if legend:
        _legend_layout(ch)
    return ch, ChartGeom(x, y, w, h, *plot, len(cats), vmax)


def scatter_chart(slide, x, y, w, h, points, font, *, color="#003087", highlight=None, highlight_names=(),
                  log_x=False, x_range=None, y_range=None, x_title="", y_title="", label_pos=None,
                  text="#1F2328", muted="#57606A", grid="#D0D7DE", y_format='0"%"'):
    """Native XY scatter. points: [(label, x, y)]. label_pos: {label: XL_LABEL_POSITION}
    to keep neighbouring labels apart. Returns the chart."""
    xy = XyChartData()
    ser = xy.add_series("points")
    for _, px, py in points:
        ser.add_data_point(px, py)
    ch = slide.shapes.add_chart(XL_CHART_TYPE.XY_SCATTER, Inches(x), Inches(y), Inches(w), Inches(h), xy).chart
    ch.font.name, ch.font.size, ch.font.color.rgb = font, Pt(14), rgb(text)
    ch.has_legend = ch.has_title = False
    va, ca = ch.value_axis, ch.category_axis
    if y_range:
        va.minimum_scale, va.maximum_scale = y_range
    if x_range:
        ca.minimum_scale, ca.maximum_scale = x_range
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = rgb(grid)
    ca.has_major_gridlines = False
    va.tick_labels.number_format, va.tick_labels.number_format_is_linked = y_format, False
    for ax, title in ((ca, x_title), (va, y_title)):
        ax.tick_labels.font.size = Pt(12)
        if title:
            ax.has_title = True
            ax.axis_title.text_frame.text = title
            r = ax.axis_title.text_frame.paragraphs[0].runs[0]
            r.font.size, r.font.name, r.font.bold, r.font.color.rgb = Pt(14), font, False, rgb(muted)
    if log_x:
        sc = ca._element.find(qn("c:scaling"))
        for e in sc.findall(qn("c:logBase")):
            sc.remove(e)
        lb = etree.Element(qn("c:logBase"))
        lb.set("val", "10")
        sc.insert(0, lb)
    sr = ch.plots[0].series[0]
    sr.format.line.fill.background()
    sr.marker.style, sr.marker.size = 8, 10
    sr.marker.format.fill.solid()
    sr.marker.format.fill.fore_color.rgb = rgb(color)
    sr.marker.format.line.fill.background()
    for i, (name, _, _) in enumerate(points):
        pt = sr.points[i]
        if highlight and name in highlight_names:
            pt.marker.format.fill.solid()
            pt.marker.format.fill.fore_color.rgb = rgb(highlight)
            pt.marker.style, pt.marker.size = 8, 12
        dl = pt.data_label
        dl.has_text_frame = True
        dl.text_frame.text = name
        dl.position = (label_pos or {}).get(name, XL_LABEL_POSITION.ABOVE)
        for r in dl.text_frame.paragraphs[0].runs:
            r.font.size, r.font.name = Pt(11), font
    return ch


# ----------------------------------------------------------- big-type slides

def set_background(slide, color):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(color)


def add_closing(prs, font, *, brand="#003087", text="Thank you", logo=None, empty_layout=None,
                notes="Thank the audience and open the floor for questions."):
    """The always-last Thank-you slide. Uses the template's own closing layout if it
    has one (fixed design left untouched, empty text placeholders removed);
    otherwise a styled slide on `empty_layout` (a layout name, default the first
    layout with no placeholders): brand background, white 56 pt text, optional logo."""
    lay = find_closing_layout(prs)
    if lay is not None:
        s = prs.slides.add_slide(lay)
        scrub_placeholders(s)
        s.notes_slide.notes_text_frame.text = notes
        return s
    lays = layouts(prs)
    if empty_layout:
        lay = lays[empty_layout]
    else:
        lay = next((l for l in prs.slide_layouts if len(l.placeholders) == 0), prs.slide_layouts[-1])
    s = prs.slides.add_slide(lay)
    for p in list(s.placeholders):
        p._element.getparent().remove(p._element)
    set_background(s, brand)
    W, H = prs.slide_width / EMU, prs.slide_height / EMU
    textbox(s, 0.9, H * 0.38, W - 2 * 0.9 - (3.6 if logo else 0), 1.4, [(text, {"size": 56, "bold": True})],
            font, color="#FFFFFF")
    if logo:
        add_picture_fit(s, logo, W - 4.2, H * 0.3, 3.4, 2.4)
    s.notes_slide.notes_text_frame.text = notes
    return s


# --------------------------------------------------- deck-level edit helpers

def delete_slide(prs, index):
    ids = prs.slides._sldIdLst
    sld = list(ids)[index]
    prs.part.drop_rel(sld.rId)
    ids.remove(sld)


def move_slide(prs, old_index, new_index):
    ids = prs.slides._sldIdLst
    el = list(ids)[old_index]
    ids.remove(el)
    ids.insert(new_index, el)


def set_hidden(slide, hidden=True):
    """Hide a slide in the slideshow (it stays in the file)."""
    if hidden:
        slide._element.set("show", "0")
    elif "show" in slide._element.attrib:
        del slide._element.attrib["show"]


def hidden_slides(prs):
    """1-based numbers of hidden slides. Exports (PDF) skip them, so page numbers
    after a hidden slide are off by one."""
    return [i + 1 for i, s in enumerate(prs.slides) if s._element.get("show") == "0"]


def find_slide(prs, title_part):
    """0-based index of the first slide whose title/first text contains the string."""
    t = title_part.lower()
    for i, s in enumerate(prs.slides):
        for sh in s.shapes:
            if sh.has_text_frame and t in sh.text_frame.text.lower():
                return i
    return None


def slide_fingerprints(prs):
    """One hash per slide: XML, notes, and the content of related parts (charts, images)."""
    out = []
    for s in prs.slides:
        h = hashlib.sha1(etree.tostring(s._element))
        if s.has_notes_slide:
            h.update(s.notes_slide.notes_text_frame.text.encode())
        for rel in s.part.rels.values():
            if rel.is_external:
                continue
            try:
                h.update(rel.target_part.blob)
            except Exception:
                pass
        out.append(h.hexdigest())
    return out
