#!/usr/bin/env python3
"""Anti-slop linter for generated .pptx decks (pitch-deck and td-deck skills).

Enforces the rules in deck-style.md that can be checked mechanically.

Usage:
    python3 lint_deck.py DECK.pptx [--template] [--max-words N] [--allow TERM ...]

Prints one ``slide N: ERROR|WARN ...`` line per finding and exits 1 if any
ERROR was found, 0 otherwise. ``--template`` skips the visual checks, which
would otherwise flag shapes inherited from an organisation's template.
"""

import argparse
import re
import sys

try:
    from pptx import Presentation
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Emu
except ImportError:
    sys.exit("lint_deck.py needs python-pptx: pip install python-pptx")

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"

BANNED = [
    r"seamless(ly)?", r"robust(ly|ness)?", r"cutting[- ]edge", r"state[- ]of[- ]the[- ]art",
    r"revolutionar(y|ize|izes|izing)", r"game[- ]chang(er|ers|ing)", r"leverag(e|es|ed|ing)",
    r"unlock(s|ed|ing)?", r"empower(s|ed|ing|ment)?", r"harness(es|ed|ing)?",
    r"delv(e|es|ed|ing)", r"elevat(e|es|ed|ing)", r"streamlin(e|es|ed|ing)",
    r"supercharg(e|es|ed|ing)", r"next[- ]gen(eration)?", r"world[- ]class", r"best[- ]in[- ]class",
    r"holistic(ally)?", r"synerg(y|ies|istic)", r"paradigm shift", r"transformative",
    r"innovative", r"powerful", r"effortless(ly)?", r"unparalleled", r"unprecedented",
    r"comprehensive", r"journey", r"landscape", r"realm", r"tapestry", r"testament",
    r"beacon", r"all[- ]in[- ]one", r"fast[- ]paced", r"ever[- ](evolving|changing)",
    r"at the heart of", r"in today's", r"it's worth noting", r"dive into", r"look no further",
    r"navigat(e|es|ing) the complexit(y|ies)", r"build(s|ing)? the future",
]
PHRASE_WARN = [r"not just\b", r"more than just\b", r"key takeaways", r"in this slide"]
LEFTOVER = [r"click to add", r"lorem ipsum", r"\bTODO\b", r"<fill in", r"\bTBD\b", r"\bXXX\b"]
EMOJI = re.compile("[\U0001F300-\U0001FAFF☀-➿⭐⬆↔-⇿]")
SHORT_WORDS = {"a", "an", "the", "of", "to", "in", "on", "at", "by", "for", "or", "and"}
FILLER_TITLES = re.compile(r"^(thank you|thanks|questions\??|q\s*&\s*a|any questions)\W*$", re.I)
PLACEHOLDER_LABEL = "DIAGRAM PLACEHOLDER"
PAGE_COUNTER = re.compile(r"^\s*(\d{1,3}\s*(/|of)\s*\d{1,3}|0\d)\s*$", re.I)
GRID = 228600  # 0.25 in, for comparing slide layouts


class Report:
    def __init__(self):
        self.lines, self.errors, self.warnings = [], 0, 0

    def error(self, n, msg):
        self.errors += 1
        self.lines.append(f"slide {n}: ERROR {msg}")

    def warn(self, n, msg):
        self.warnings += 1
        self.lines.append(f"slide {n}: WARN  {msg}")


def shape_text(shape):
    if shape.has_text_frame:
        return shape.text_frame.text
    if getattr(shape, "has_table", False) and shape.has_table:
        return "\n".join(c.text for row in shape.table.rows for c in row.cells)
    return ""


def iter_shapes(shapes):
    for sh in shapes:
        if sh.shape_type == 6:  # group
            yield from iter_shapes(sh.shapes)
        else:
            yield sh


def max_font_pt(shape):
    sizes = [r.font.size.pt for p in shape.text_frame.paragraphs for r in p.runs if r.font.size]
    return max(sizes) if sizes else None


def find_title(slide):
    """Title placeholder if it has text, else the topmost text box of >= 24 pt in the top 40%
    (so a big number below the title isn't mistaken for it), else the largest font there."""
    t = slide.shapes.title
    if t is not None and t.text_frame.text.strip():
        return t
    height = slide.part.package.presentation_part.presentation.slide_height
    cands = [sh for sh in iter_shapes(slide.shapes)
             if sh.has_text_frame and sh.text_frame.text.strip() and sh.top is not None
             and sh.top <= height * 0.4 and PLACEHOLDER_LABEL not in sh.text_frame.text]
    big = [sh for sh in cands if (max_font_pt(sh) or 0) >= 24]
    if big:
        return min(big, key=lambda sh: sh.top)
    return max(cands, key=lambda sh: max_font_pt(sh) or 0, default=None)


def is_title(sh, title):
    # python-pptx builds new proxy objects per iteration, so compare ids, not identity
    return title is not None and sh.shape_id == title.shape_id


def words(text):
    return re.findall(r"[A-Za-z0-9][\w'’.-]*", text)


def check_copy(n, slide, title, rep, banned_re, max_words, is_first):
    title_text = title.text_frame.text.strip() if title is not None else ""
    notes = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""
    texts = [shape_text(sh) for sh in iter_shapes(slide.shapes)]
    everything = "\n".join(texts + [notes])

    if "—" in everything:
        rep.error(n, "em dash (—) found; use an en dash for ranges, otherwise rephrase")
    for m in sorted({m.group(0).lower() for m in banned_re.finditer(everything)}):
        rep.error(n, f"banned vocabulary: '{m}'")
    for pat in PHRASE_WARN:
        if re.search(pat, everything, re.I):
            rep.warn(n, f"slop phrase pattern: /{pat}/")
    for pat in LEFTOVER:
        if re.search(pat, everything, re.I):
            rep.error(n, f"leftover placeholder text matching /{pat}/")
    if EMOJI.search(everything):
        rep.error(n, "emoji or decorative glyph found")
    if "\u00b7" in everything:
        rep.warn(n, "middot (·) used as a separator")
    if not notes.strip():
        rep.error(n, "no speaker notes")

    if title_text:
        lines = [l.strip() for l in title_text.splitlines() if l.strip()]
        if lines and lines[-1][-1] in ".!?":
            rep.error(n, f"title ends with '{lines[-1][-1]}': {title_text!r}")
        letters = [c for c in title_text if c.isalpha()]
        if len(words(title_text)) >= 2 and letters and all(c.isupper() for c in letters):
            rep.error(n, f"title in ALL CAPS: {title_text!r}")
        tail = [w for w in words(title_text)[1:] if len(w) >= 4]
        if len(tail) >= 3 and all(w[0].isupper() for w in tail):
            rep.warn(n, f"title looks like Title Case (fine only if these are proper nouns): {title_text!r}")
        for l in lines[:-1]:
            last = l.split()[-1].lower() if l.split() else ""
            if last in SHORT_WORDS:
                rep.warn(n, f"title line ends with short word '{last}'")
        if re.search(r"\(\s*\d+\s*/\s*\d+\s*\)|\b\d+\s*/\s*\d+\s*$", title_text):
            rep.error(n, f"series counter in title: {title_text!r}")
        if not is_first and FILLER_TITLES.match(title_text):
            rep.warn(n, f"filler closing slide: {title_text!r}")

    height = slide.part.package.presentation_part.presentation.slide_height
    in_footer = lambda sh: sh.top is not None and sh.top >= height * 0.92  # footer band: chrome, not body
    body = "\n".join(shape_text(sh) for sh in iter_shapes(slide.shapes)
                     if not is_title(sh, title) and not in_footer(sh))
    body = body.replace(PLACEHOLDER_LABEL, "")
    wc = len(words(body))
    if not is_first and wc > max_words:
        rep.warn(n, f"{wc} body words (limit {max_words}); consider splitting")
    bullets = max((sum(1 for p in sh.text_frame.paragraphs if p.text.strip())
                   for sh in iter_shapes(slide.shapes)
                   if not is_title(sh, title) and sh.has_text_frame and PLACEHOLDER_LABEL not in sh.text_frame.text),
                  default=0)
    return title_text, bullets


def layout_signature(slide):
    """Shape geometry snapped to a 0.25 in grid: equal signatures = same layout."""
    return tuple(sorted(
        (sh.left // GRID, sh.top // GRID, sh.width // GRID, sh.height // GRID)
        for sh in iter_shapes(slide.shapes) if None not in (sh.left, sh.top, sh.width, sh.height)))


def runs(values, pred, min_len):
    """Yield (start, end) 1-based slide ranges where pred holds for >= min_len consecutive items."""
    start = None
    for i, v in enumerate(values + [None], start=1):
        if v is not None and pred(i, v):
            start = start or i
        else:
            if start and i - start >= min_len:
                yield start, i - 1
            start = None


def spPr_fill(sh):
    spPr = sh._element.find(f"{P}spPr")
    if spPr is None:
        return None
    for tag in ("noFill", "solidFill", "gradFill", "pattFill", "blipFill"):
        if spPr.find(f"{A}{tag}") is not None:
            return tag
    return None


def is_filled(sh):
    fill = spPr_fill(sh)
    if fill:
        return fill != "noFill"
    style = sh._element.find(f"{P}style")
    if style is not None:
        ref = style.find(f"{A}fillRef")
        return ref is not None and ref.get("idx", "0") != "0"
    return False


def has_shadow(sh):
    # A theme effectRef counts even when spPr carries an empty effectLst override:
    # LibreOffice still renders the theme shadow in that case.
    el = sh._element
    spPr = el.find(f"{P}spPr")
    if spPr is not None:
        eff = spPr.find(f"{A}effectLst")
        if eff is not None and len(eff) > 0:
            return True
    style = el.find(f"{P}style")
    if style is not None:
        ref = style.find(f"{A}effectRef")
        return ref is not None and ref.get("idx", "0") != "0"
    return False


def check_visual(n, slide, title, prs, rep, fonts, accents, is_first):
    W, H = prs.slide_width, prs.slide_height
    bold_runs = 0
    for sh in iter_shapes(slide.shapes):
        el = sh._element
        geom = el.find(f".//{A}prstGeom")
        prst = geom.get("prst") if geom is not None else ""
        if "round" in prst.lower() or prst.lower().startswith("snip"):
            rep.error(n, f"rounded/snipped shape '{sh.name}' ({prst})")
        if el.find(f".//{A}gradFill") is not None:
            rep.error(n, f"gradient fill on '{sh.name}'")
        if has_shadow(sh):
            rep.error(n, f"shadow/effect on '{sh.name}' (use add_textbox, or clear the effect)")
        if None not in (sh.left, sh.top, sh.width, sh.height):
            tol = Emu(12700)  # 1 pt
            if sh.left < -tol or sh.top < -tol or sh.left + sh.width > W + tol or sh.top + sh.height > H + tol:
                rep.error(n, f"'{sh.name}' extends beyond the slide")

        if (not (sh.has_text_frame and sh.text_frame.text.strip()) and is_filled(sh)
                and None not in (sh.width, sh.height) and min(sh.width, sh.height) > 0
                and min(sh.width, sh.height) < Emu(73152)
                and max(sh.width, sh.height) / min(sh.width, sh.height) > 8):
            rep.error(n, f"decorative accent bar '{sh.name}' (thin filled shape)")
        if not sh.has_text_frame or not sh.text_frame.text.strip():
            if sh.is_placeholder and sh.has_text_frame:
                rep.warn(n, f"empty placeholder '{sh.name}' (remove it or use the Blank layout)")
            continue
        text = sh.text_frame.text
        if PAGE_COUNTER.match(text):
            rep.error(n, f"page-counter / decorative number chrome: {text.strip()!r}")
        if is_title(sh, title):
            colours = set()
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    try:
                        if r.text.strip() and r.font.color and r.font.color.type is not None:
                            colours.add(str(r.font.color.rgb))
                    except AttributeError:
                        pass
            if len(colours) > 1:
                rep.error(n, f"title uses {len(colours)} colours (accent on part of a title)")
        if is_filled(sh) and sh.shape_type != 17:  # 17 = text box without fill
            rep.error(n, f"text inside filled shape '{sh.name}': {text.strip()[:40]!r} (badge/pill/card)")
        if (title is not None and not is_title(sh, title) and sh.top is not None and title.top is not None
                and sh.top + sh.height <= title.top + Emu(38100) and len(words(text)) <= 6):
            rep.error(n, f"eyebrow/kicker above the title: {text.strip()!r}")
        for p in sh.text_frame.paragraphs:
            if (not is_first and not is_title(sh, title) and p.alignment == PP_ALIGN.CENTER
                    and p.text.strip() and PLACEHOLDER_LABEL not in text):
                rep.warn(n, f"centred body text: {p.text.strip()[:40]!r}")
            for r in p.runs:
                if r.font.name:
                    fonts.add(r.font.name)
                if r.font.italic:
                    rep.warn(n, f"italic run: {r.text.strip()[:30]!r}")
                if r.font.bold and not is_title(sh, title) and r.text.strip():
                    bold_runs += 1
                try:
                    rgb = r.font.color.rgb if r.font.color and r.font.color.type is not None else None
                except AttributeError:
                    rgb = None
                if rgb is not None:
                    rr, gg, bb = rgb[0], rgb[1], rgb[2]
                    if max(rr, gg, bb) - min(rr, gg, bb) > 40:
                        accents.add(str(rgb))
    if bold_runs > 1:
        rep.warn(n, f"{bold_runs} bold runs in body (bold at most one key term/number per slide)")


def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[1])
    ap.add_argument("deck")
    ap.add_argument("--template", action="store_true", help="deck is built on a template: copy checks only")
    ap.add_argument("--max-words", type=int, default=60)
    ap.add_argument("--allow", nargs="*", default=[], help="canonical terms exempt from the vocabulary ban")
    args = ap.parse_args()

    prs = Presentation(args.deck)
    allow = {a.lower() for a in args.allow}
    banned = [b for b in BANNED if not any(re.fullmatch(b, a, re.I) for a in allow)]
    banned_re = re.compile(r"\b(" + "|".join(banned) + r")\b", re.I)

    rep, fonts, accents, titles, bullets, layouts = Report(), set(), set(), [], [], []
    for n, slide in enumerate(prs.slides, start=1):
        title = find_title(slide)
        t, b = check_copy(n, slide, title, rep, banned_re, args.max_words, n == 1)
        titles.append(t)
        bullets.append(b)
        layouts.append(layout_signature(slide))
        if not args.template:
            check_visual(n, slide, title, prs, rep, fonts, accents, n == 1)

    for a, b in runs(bullets, lambda i, v: i > 1 and 3 <= v <= 4, 3):
        rep.warn(a, f"slides {a}-{b} all use 3-4 bullets; vary the form (sentence, number, table, diagram)")
    if not args.template:
        same = lambda i, v: i > 2 and v and v == layouts[i - 2]
        for a, b in runs(layouts, same, 3):
            rep.warn(a - 1, f"slides {a - 1}-{b} share an identical layout; no more than 3 in a row")

    seen = {}
    for n, t in enumerate(titles, start=1):
        key = t.strip().lower()
        if key and key in seen:
            rep.error(n, f"duplicate title (also slide {seen[key]}): {t!r}")
        seen.setdefault(key, n)
    if not args.template:
        if len(fonts) > 2:
            rep.error(0, f"{len(fonts)} font families used: {sorted(fonts)} (max 2, ideally 1)")
        if len(accents) > 2:
            rep.warn(0, f"{len(accents)} saturated text colours: {sorted(accents)} (use one accent)")

    for line in rep.lines:
        print(line.replace("slide 0:", "deck:"))
    print(f"{rep.errors} error(s), {rep.warnings} warning(s)"
          + (" [template mode: visual checks skipped]" if args.template else ""))
    sys.exit(1 if rep.errors else 0)


if __name__ == "__main__":
    main()
